#!/usr/bin/env python3
"""Build docs/HCM_Academy_QA_Review.xlsx: a reviewer workbook for HCM Academy.
Sheets: How to use, Overall, Sections, Courses & assets, Assessments, Issue log.
Run: python3 tools/build_qa_workbook.py"""
import html, json, os, re, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
from course_content import COURSES, VERIFY  # noqa: E402
ROOT = os.path.dirname(HERE)
SITE = "https://me5231979.github.io/HCM_Academy/"
BLACK, GOLD, OAK, CREAM, SAND = "1C1C1C", "CFAE70", "946E24", "F5F3EF", "E0D5C0"

# landing structure: topics, sections, rows, status
landing = open(os.path.join(ROOT, "index.html"), encoding="utf-8").read()
TOPICS, ROWS = [], []
for m in re.finditer(r'<section class="card[^"]*" id="([a-z]+)" data-short="([^"]*)">(.*?)</section>', landing, re.S):
    cid, short, body = m.groups()
    name = html.unescape(re.search(r"<h3>(.*?)</h3>", body).group(1)); no = re.search(r'<span class="no"[^>]*>(\d+)</span>', body).group(1)
    sub = ""; live = soon = 0
    for part in re.split(r'(<h4 class="subhead">.*?</h4>)', body):
        h = re.match(r'<h4 class="subhead">(.*?)</h4>', part)
        if h: sub = html.unescape(h.group(1)); continue
        for li in re.finditer(r'<li([^>]*)><span class="nm"><span class="tt">(.*?)</span>', part):
            attrs, title = li.group(1), html.unescape(li.group(2))
            if "data-see" in attrs or "data-assess" in attrs: continue
            slug = re.search(r'data-course="([^"]+)"', attrs)
            ROWS.append(dict(no=no, topic=name, section=sub or name, title=title, slug=slug.group(1) if slug else None))
            if slug: live += 1
            else: soon += 1
    TOPICS.append(dict(no=no, name=name, live=live, soon=soon))
BY = {c["slug"]: c for c in COURSES}

def source_of(c):
    p = c.get("pending", "")
    if p.startswith("Built from Oracle"): return "Oracle docs and official sources (pending SME review)"
    return "Vanderbilt SME sessions, transcripts, and guides"

def open_items(c):
    items = [w.replace("Open item: ", "") for w in c["watch"] if w.startswith("Open item")]
    items += VERIFY.get(c["slug"], [])
    seen, out = set(), []
    for i in items:
        k = i.lower().strip(" .")
        if not any(k in x or x in k for x in seen): seen.add(k); out.append(i[0].upper() + i[1:])
    return "\n".join(f"- {i}" for i in out)

# assessments
t = open(os.path.join(ROOT, "assets", "assessments.js"), encoding="utf-8").read()
topics_js = dict(re.findall(r"(\w+): '([^']+)'", t[t.index("topics:"):t.index("pre:")]))
def qs(kind):
    seg = t[t.index(kind + ": ["):]; seg = seg[:seg.index("\n  ]")]
    return [json.loads(l.strip().rstrip(",")) for l in seg.splitlines()[1:] if l.strip().startswith("{")]

wb = Workbook()
thin = Side(style="thin", color=SAND)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill("solid", fgColor=BLACK); HEAD_FONT = Font(name="Arial", bold=True, color=GOLD, size=10)
INPUT_FILL = PatternFill("solid", fgColor="FFFDF7"); PRE_FILL = PatternFill("solid", fgColor=CREAM)
WRAP = Alignment(wrap_text=True, vertical="top")

def sheet(title, cols, rows, input_cols, lists):
    ws = wb.create_sheet(title)
    ws.append([c[0] for c in cols])
    for i, (name, width) in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i); cell.fill = HEAD_FILL; cell.font = HEAD_FONT; cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[1].height = 32
    for r in rows: ws.append(r)
    n = ws.max_row
    for row in ws.iter_rows(min_row=2, max_row=max(n, 2)):
        for cell in row:
            cell.alignment = WRAP; cell.border = BORDER; cell.font = Font(name="Arial", size=10)
            hdr = cols[cell.column - 1][0]
            cell.fill = INPUT_FILL if hdr in input_cols else PRE_FILL
    ws.freeze_panes = "B2" if title != "How to use" else None
    ws.auto_filter.ref = ws.dimensions
    last = max(n, 200)
    for hdr, ref in lists.items():
        col = [c[0] for c in cols].index(hdr) + 1; letter = ws.cell(row=1, column=col).column_letter
        dv = DataValidation(type="list", formula1=f"=Lists!${ref}", allow_blank=True); dv.error = "Pick a value from the list."; dv.showErrorMessage = True
        ws.add_data_validation(dv); dv.add(f"{letter}2:{letter}{last}")
        if hdr == "Status":
            rng = f"{letter}2:{letter}{last}"
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Done"'], fill=PatternFill("solid", fgColor="E3F0E4")))
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Open"'], fill=PatternFill("solid", fgColor="FBEFD5")))
        if hdr == "Priority":
            ws.conditional_formatting.add(f"{letter}2:{letter}{last}", CellIsRule(operator="equal", formula=['"High"'], font=Font(name="Arial", bold=True, color="9C2B2B")))
    # input area styling for empty rows too
    for r in range(n + 1, last + 1):
        for i, (hdr, _) in enumerate(cols, 1):
            c = ws.cell(row=r, column=i); c.border = BORDER; c.alignment = WRAP
            c.fill = INPUT_FILL if hdr in input_cols else PatternFill(fill_type=None)
    return ws

# Lists (hidden)
L = wb.active; L.title = "Lists"
lists = {"A": ["5 Excellent", "4 Good", "3 OK", "2 Needs work", "1 Not usable"], "B": ["High", "Medium", "Low"],
         "C": ["Open", "In progress", "Done", "Won't change"], "D": ["Accurate", "Needs change", "Wrong", "Not reviewed"],
         "E": ["Approve", "Edit", "Remove"], "F": ["Content", "Steps or screens", "Quiz", "Job aid", "Video", "Design", "Accessibility", "Navigation", "Typo", "Missing topic", "Other"]}
for col, vals in lists.items():
    for i, v in enumerate(vals, 1): L[f"{col}{i}"] = v
REF = {k: f"{k}$1:${k}${len(v)}" for k, v in lists.items()}

# How to use
H = wb.create_sheet("How to use", 0)
H.column_dimensions["A"].width = 28; H.column_dimensions["B"].width = 100
H["A1"] = "HCM Academy QA review"; H["A1"].font = Font(name="Georgia", size=22, color=BLACK)
H["A2"] = "V.01 draft review workbook. Add your feedback in the cream input columns; gray columns are pre-filled."; H["A2"].font = Font(name="Arial", size=10, italic=True, color="5E5B55")
info = [("Site", SITE), ("HCM Playbook (PDF)", SITE + "playbook/HCM_Playbook.pdf"), ("Return to", "pcb@vanderbilt.edu"),
        ("Reviewer name", ""), ("Review date", ""), ("",""),
        ("Overall", "Feedback on the Academy as a whole: landing page, navigation, assessments, locking, playbook, design, accessibility."),
        ("Sections", "One row per topic (00 to 07): is anything missing, out of order, or in the wrong section?"),
        ("Courses & assets", "One row per course. Rate the walkthrough steps, job aid, knowledge check, and video or detailed job aid. 'Open items to confirm' lists what we already know needs a Vanderbilt check."),
        ("Assessments", "Every pre-assessment and final assessment question with its correct answer. Approve, Edit, or Remove each one."),
        ("Issue log", "Anything else: bugs, typos, broken links, or ideas. One row per issue."), ("",""),
        ("Rating", "5 Excellent, 4 Good, 3 OK, 2 Needs work, 1 Not usable."),
        ("Priority", "High: wrong or blocks someone doing the work. Medium: unclear or incomplete. Low: polish."),
        ("Status", "Open, In progress, Done, Won't change. PCB updates Status and Response."),
        ("Source", "Courses marked 'pending SME review' were built from Oracle documentation and official sources. They need a closer look for Vanderbilt accuracy.")]
for i, (k, v) in enumerate(info, 4):
    H[f"A{i}"] = k; H[f"B{i}"] = v
    H[f"A{i}"].font = Font(name="Arial", bold=True, size=10, color=OAK); H[f"B{i}"].font = Font(name="Arial", size=10); H[f"B{i}"].alignment = Alignment(wrap_text=True, vertical="top")
    if k in ("Reviewer name", "Review date"): H[f"B{i}"].fill = INPUT_FILL; H[f"B{i}"].border = BORDER
    if v.startswith("http"): H[f"B{i}"].hyperlink = v; H[f"B{i}"].font = Font(name="Arial", size=10, color="1F4E8C", underline="single")

# Overall
areas = [("Landing page", "Is it clear what the Academy is and where to start?"), ("Topic tabs and order", "Are the eight topics the right ones, in the right order?"),
         ("Search", "Does search find what you'd look for?"), ("Pre-assessment", "Right questions? Do results point to the right topics?"),
         ("Lock until pre-assessment", "Is locking every course until the pre-assessment the right experience?"), ("Progress and completion", "Do checks and counts behave as expected?"),
         ("Final assessment", "Right coverage and difficulty? Is 90% the right pass mark?"), ("Course format", "Intro, steps, 3-question check, job aid: right length and depth?"),
         ("Job aids", "Would you print and use them? Anything missing?"), ("HCM Playbook", "Cover, contents, dividers, guides: useful and on brand?"),
         ("Videos", "Is the Guidde video placement right? Which courses need videos next?"), ("Branding and design", "Does it look and read like Vanderbilt?"),
         ("Accessibility and mobile", "Usable with keyboard, screen reader, and on a phone?"), ("Terminology", "Do names match what HCMs say and see in Oracle and UKG?"),
         ("Missing topics", "What should be added that isn't here?"), ("Other", "")]
sheet("Overall", [("Area", 24), ("What to review", 44), ("Reviewer", 16), ("Rating", 14), ("Feedback or change requested", 60), ("Priority", 11), ("Status", 13), ("PCB response", 40)],
      [[a, q, "", "", "", "", "Open", ""] for a, q in areas], {"Reviewer", "Rating", "Feedback or change requested", "Priority"}, {"Rating": REF["A"], "Priority": REF["B"], "Status": REF["C"]})

# Sections
sheet("Sections", [("No.", 6), ("Topic", 34), ("Courses live", 9), ("Coming soon", 9), ("Reviewer", 16), ("Rating", 14), ("Missing or misplaced topics", 44), ("Feedback or change requested", 54), ("Priority", 11), ("Status", 13), ("PCB response", 36)],
      [[t["no"], t["name"], t["live"], t["soon"], "", "", "", "", "", "Open", ""] for t in TOPICS],
      {"Reviewer", "Rating", "Missing or misplaced topics", "Feedback or change requested", "Priority"}, {"Rating": REF["A"], "Priority": REF["B"], "Status": REF["C"]})

# Courses & assets
crows = []
for i, r in enumerate(ROWS, 1):
    c = BY.get(r["slug"]) if r["slug"] else None
    if c:
        extra = "Video; detailed job aid" if c.get("video") or c.get("detailed") else ""
        crows.append([f"C{i:03d}", f'{r["no"]} {r["topic"]}', r["section"], c["name"], "Live", source_of(c), SITE + f"courses/{c['slug']}/", c["system"], c["time"], extra,
                      open_items(c), "", "", "", "", "", "", "", "", "Open", ""])
    else:
        crows.append([f"C{i:03d}", f'{r["no"]} {r["topic"]}', r["section"], r["title"], "Coming soon", "No content yet", "", "", "", "", "Needs a source: SME session, guide, or recording.", "", "", "", "", "", "", "", "", "Open", ""])
for name, url, kind in [("Pre-assessment", SITE + "assessment/pre/", "Assessment"), ("Final assessment", SITE + "assessment/final/", "Assessment"), ("HCM Playbook (PDF)", SITE + "playbook/HCM_Playbook.pdf", "Print book")]:
    crows.append([f"A{len(crows)+1:03d}", "All topics", kind, name, "Live", "PCB", url, kind, "", "", "", "", "", "", "", "", "", "", "", "Open", ""])
ws = sheet("Courses & assets",
      [("ID", 7), ("Topic", 22), ("Section", 18), ("Course or asset", 34), ("Course status", 12), ("Source", 24), ("Link", 26), ("System", 14), ("Time", 8), ("Extras", 14),
       ("Open items to confirm", 46), ("Reviewer", 14), ("Walkthrough steps", 14), ("Job aid", 13), ("Knowledge check", 14), ("Video or detailed aid", 14),
       ("Feedback or change requested", 54), ("Issue type", 15), ("Priority", 10), ("Status", 12), ("PCB response", 34)],
      crows, {"Reviewer", "Walkthrough steps", "Job aid", "Knowledge check", "Video or detailed aid", "Feedback or change requested", "Issue type", "Priority"},
      {"Walkthrough steps": REF["D"], "Job aid": REF["D"], "Knowledge check": REF["D"], "Video or detailed aid": REF["D"], "Issue type": REF["F"], "Priority": REF["B"], "Status": REF["C"]})
for row in ws.iter_rows(min_row=2, max_row=len(crows) + 1, min_col=7, max_col=7):
    for cell in row:
        if cell.value: cell.hyperlink = cell.value; cell.value = "Open"; cell.font = Font(name="Arial", size=10, color="1F4E8C", underline="single")
st_col = [c.value for c in ws[1]].index("Course status") + 1
L_ = ws.cell(row=1, column=st_col).column_letter
ws.conditional_formatting.add(f"{L_}2:{L_}{len(crows)+1}", CellIsRule(operator="equal", formula=['"Coming soon"'], font=Font(name="Arial", size=10, italic=True, color="8A8478")))

# Assessments
arows = []
for kind, label in (("pre", "Pre-assessment"), ("final", "Final assessment")):
    for i, q in enumerate(qs(kind), 1):
        opts = "\n".join(f"{'ABC'[j]}) {o}" for j, o in enumerate(q["o"]))
        arows.append([label, i, topics_js.get(q["topic"], q["topic"]), q["q"], opts, f"{'ABC'[q['a']]}) {q['o'][q['a']]}", q.get("why", ""), "Yes" if q.get("draft") else "", "", "", "", "", "Open", ""])
sheet("Assessments", [("Assessment", 15), ("Q", 5), ("Topic", 26), ("Question", 44), ("Options", 40), ("Correct answer", 30), ("Feedback shown", 30), ("Draft (SME to confirm)", 11),
      ("Reviewer", 14), ("Decision", 12), ("Suggested edit", 46), ("Priority", 10), ("Status", 12), ("PCB response", 30)],
      arows, {"Reviewer", "Decision", "Suggested edit", "Priority"}, {"Decision": REF["E"], "Priority": REF["B"], "Status": REF["C"]})

# Issue log
sheet("Issue log", [("Date", 12), ("Reviewer", 16), ("Page or asset", 30), ("Link", 30), ("Issue type", 15), ("What's wrong or what to change", 60), ("Screenshot link", 24), ("Priority", 10), ("Status", 12), ("PCB response", 36)],
      [], {"Date", "Reviewer", "Page or asset", "Link", "Issue type", "What's wrong or what to change", "Screenshot link", "Priority"}, {"Issue type": REF["F"], "Priority": REF["B"], "Status": REF["C"]})

L.sheet_state = "hidden"
wb.move_sheet("Lists", offset=len(wb.sheetnames))
for ws in wb.worksheets:
    ws.sheet_properties.tabColor = GOLD if ws.title != "How to use" else BLACK
    ws.sheet_view.zoomScale = 100
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.print_title_rows = "1:1" if ws.title != "How to use" else None
os.makedirs(os.path.join(ROOT, "docs"), exist_ok=True)
out = os.path.join(ROOT, "docs", "HCM_Academy_QA_Review.xlsx")
wb.save(out)
print(f"wrote {out}: {len(crows)} courses and assets, {len(arows)} assessment questions, {len(TOPICS)} sections")
